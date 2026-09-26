import zipfile
from pathlib import Path

import docx
import openpyxl
import pptx
from pptx.util import Inches


def make_docx(path: Path) -> Path:
    document = docx.Document()
    document.add_heading("KeepReadable fixture", level=1)
    document.add_paragraph("A deterministic paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    document.save(path)
    return path


def make_xlsx(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Value"])
    sheet.append(["alpha", 1])
    workbook.create_sheet("Second")
    workbook.save(path)
    workbook.close()
    return path


def make_pptx(path: Path) -> Path:
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "KeepReadable fixture"
    presentation.save(path)
    return path


def _rewrite(src: Path, dst: Path, replacements: dict[str, bytes | None]) -> Path:
    with zipfile.ZipFile(src) as source, zipfile.ZipFile(dst, "w") as target:
        for info in source.infolist():
            replacement = replacements.get(info.filename, source.read(info.filename))
            if replacement is not None:
                target.writestr(info, replacement)
    return dst


def remove_zip_member(src: Path, dst: Path, name: str) -> Path:
    return _rewrite(src, dst, {name: None})


def corrupt_zip_member_xml(src: Path, dst: Path, name: str) -> Path:
    return _rewrite(src, dst, {name: b"<broken"})


def make_fake_ole(path: Path) -> Path:
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 4096)
    return path
