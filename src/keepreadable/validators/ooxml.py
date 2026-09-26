import zipfile
from pathlib import Path
from typing import Any, Literal

import docx
import openpyxl
import pptx
from defusedxml import ElementTree

from keepreadable.domain.enums import CheckStatus, SupportTier
from keepreadable.utilities.cancellation import CancellationToken
from keepreadable.utilities.filesystem import extended_path
from keepreadable.validators.base import (
    EvidenceCode,
    ValidationDepth,
    ValidationResult,
    guard_validation,
)

_KIND_MAIN_PART = {
    "docx": "word/document.xml",
    "xlsx": "xl/workbook.xml",
    "pptx": "ppt/presentation.xml",
}
_OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")
_MAX_XML_PART = 50 * 1024 * 1024


class OoxmlValidator:
    tier = SupportTier.STRUCTURAL

    def __init__(self, kind: Literal["docx", "xlsx", "pptx"]) -> None:
        self.kind = kind
        self.name = f"ooxml-{kind}"

    def version(self) -> str:
        return "1"

    @guard_validation
    def validate(
        self,
        path: Path,
        *,
        depth: ValidationDepth,
        cancel: CancellationToken | None = None,
    ) -> ValidationResult:
        if cancel is not None:
            cancel.raise_if_cancelled()
        inspected = self._inspect_package(path, depth, cancel)
        if isinstance(inspected, ValidationResult):
            return inspected
        details, warnings_found = inspected
        if depth is ValidationDepth.DEEP:
            if cancel is not None:
                cancel.raise_if_cancelled()
            try:
                details.update(self._load_with_library(path))
            except Exception as exc:
                return ValidationResult(
                    self.name,
                    self.version(),
                    self.tier,
                    depth,
                    CheckStatus.PASSED,
                    CheckStatus.FAILED,
                    "The OOXML application library could not read the package content.",
                    details,
                    tuple(warnings_found),
                    (repr(exc),),
                    (EvidenceCode.DECODE_FAILURE,),
                    ("required package parts", "XML parsing", "application library load"),
                    ("visual rendering / layout fidelity",),
                )
        codes = (EvidenceCode.VALIDATION_WARNING,) if warnings_found else ()
        return ValidationResult(
            self.name,
            self.version(),
            self.tier,
            depth,
            CheckStatus.PASSED,
            CheckStatus.PASSED if depth is ValidationDepth.DEEP else CheckStatus.NOT_CHECKED,
            "Structurally readable OOXML package.",
            details,
            tuple(warnings_found),
            (),
            codes,
            (
                ("required package parts", "XML parsing", "application library load")
                if depth is ValidationDepth.DEEP
                else ("required package parts", "main XML parsing")
            ),
            (
                ("visual rendering / layout fidelity",)
                if depth is ValidationDepth.DEEP
                else ("application library load", "visual rendering / layout fidelity")
            ),
        )

    def _inspect_package(
        self,
        path: Path,
        depth: ValidationDepth,
        cancel: CancellationToken | None,
    ) -> tuple[dict[str, Any], list[str]] | ValidationResult:
        try:
            archive = zipfile.ZipFile(extended_path(path))
        except zipfile.BadZipFile as exc:
            with open(extended_path(path), "rb") as stream:
                magic = stream.read(8)
            return (
                self._protected(depth)
                if magic == _OLE_MAGIC
                else self._structural_failure(depth, exc)
            )
        with archive:
            names = set(archive.namelist())
            required = {"[Content_Types].xml", "_rels/.rels", _KIND_MAIN_PART[self.kind]}
            missing = sorted(required - names)
            if missing:
                return self._structural_failure(
                    depth, ValueError(f"missing required parts: {', '.join(missing)}"), missing
                )
            for member in ("[Content_Types].xml", "_rels/.rels", _KIND_MAIN_PART[self.kind]):
                failure = self._parse_member(archive, member, depth)
                if failure is not None:
                    return failure
            details: dict[str, Any] = {"package_kind": self.kind, "part_count": len(names)}
            if depth is ValidationDepth.QUICK:
                return details, []
            return self._inspect_all_xml(archive, details, depth, cancel)

    def _inspect_all_xml(
        self,
        archive: zipfile.ZipFile,
        details: dict[str, Any],
        depth: ValidationDepth,
        cancel: CancellationToken | None,
    ) -> tuple[dict[str, Any], list[str]] | ValidationResult:
        warnings_found: list[str] = []
        for info in archive.infolist():
            if cancel is not None:
                cancel.raise_if_cancelled()
            if not info.filename.casefold().endswith((".xml", ".rels")):
                continue
            if info.file_size > _MAX_XML_PART:
                warnings_found.append(f"{info.filename}: part not parsed due to size")
                continue
            failure = self._parse_member(archive, info.filename, depth)
            if failure is not None:
                return failure
        return details, warnings_found

    def _parse_member(
        self, archive: zipfile.ZipFile, member: str, depth: ValidationDepth
    ) -> ValidationResult | None:
        try:
            ElementTree.fromstring(archive.read(member))
        except Exception as exc:
            return self._structural_failure(depth, exc, [member])
        return None

    def _load_with_library(self, path: Path) -> dict[str, Any]:
        path_value = extended_path(path)
        if self.kind == "docx":
            document = docx.Document(path_value)
            return {
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
            }
        if self.kind == "xlsx":
            workbook = openpyxl.load_workbook(path_value, read_only=True)
            try:
                names = list(workbook.sheetnames)
            finally:
                workbook.close()
            return {"sheet_count": len(names), "sheet_names": names}
        presentation = pptx.Presentation(path_value)
        return {"slide_count": len(presentation.slides)}

    def _protected(self, depth: ValidationDepth) -> ValidationResult:
        return ValidationResult(
            self.name,
            self.version(),
            self.tier,
            depth,
            CheckStatus.NOT_CHECKED,
            CheckStatus.PROTECTED,
            (
                "This package uses a binary container, which usually means it is "
                "password-protected or a legacy format; its contents were not validated."
            ),
            {"binary_container": True},
            (),
            (),
            (EvidenceCode.PROTECTED_CONTENT,),
            ("container signature",),
            ("OOXML package structure", "visual rendering / layout fidelity"),
        )

    def _structural_failure(
        self,
        depth: ValidationDepth,
        error: Exception,
        members: list[str] | None = None,
    ) -> ValidationResult:
        return ValidationResult(
            self.name,
            self.version(),
            self.tier,
            depth,
            CheckStatus.FAILED,
            CheckStatus.NOT_CHECKED,
            "OOXML package structure could not be read.",
            {"failed_members": members or []},
            (),
            (repr(error),),
            (EvidenceCode.STRUCTURAL_FAILURE,),
            ("OOXML package structure",),
            ("application library load", "visual rendering / layout fidelity"),
        )
